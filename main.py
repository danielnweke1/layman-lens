import os
import io
from fastapi import FastAPI, HTTPException, File, UploadFile
from fastapi.responses import StreamingResponse, Response
import PyPDF2
import docx
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
import openai
from dotenv import load_dotenv

# Load environment variables
load_dotenv()

# Initialize FastAPI app
app = FastAPI(title="The Layman Lens API")

# Configure CORS so the frontend can communicate with the backend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # Allow all origins for local development
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# Initialize OpenAI Client
try:
    oai_key = os.getenv("OPENAI_API_KEY")
    if oai_key and oai_key.startswith('"') and oai_key.endswith('"'):
        oai_key = oai_key[1:-1]
    openai_client = openai.OpenAI(api_key=oai_key) if oai_key else None
except Exception as e:
    print("Warning: Failed to initialize OpenAI Client. Make sure OPENAI_API_KEY is set in your environment or .env file.")
    openai_client = None

class Message(BaseModel):
    role: str
    content: str

class ChatRequest(BaseModel):
    full_context: str
    highlighted_text: str
    chat_history: list[Message] = []
    new_question: str | None = None
    complexity: str = "5-year-old"

class ExportRequest(BaseModel):
    text: str

def get_system_instruction(complexity: str):
    return f"""
You are an expert at simplifying dense, complex text filled with jargon from ANY field (medical, legal, financial, technical, academic, etc.).
Your goal is to translate the highlighted text into plain English tailored specifically for a {complexity}.

Guidelines:
1. Identify the core message of the text based on the surrounding document context.
2. BE CONCISE. Your response MUST be 20 sentences maximum. Never write long paragraphs.
3. If the highlighted portion is a single word, give a punchy, one-sentence definition based on the context.
4. Adjust your vocabulary and tone to match the comprehension level of a {complexity}.
5. If appropriate, use a very brief analogy.
6. NEVER be condescending.
7. Return ONLY the simplified explanation, formatted nicely in Markdown.
"""

@app.post("/simplify-stream")
async def simplify_text_stream(request: ChatRequest):
    if not openai_client:
        raise HTTPException(status_code=500, detail="OpenAI Client is not initialized.")
    
    if not request.highlighted_text:
        raise HTTPException(status_code=400, detail="No highlighted text provided.")

    async def generate():
        try:
            messages = []
            messages.append({"role": "system", "content": get_system_instruction(request.complexity)})
            initial_prompt = f"Here is the full document for context:\n{request.full_context}\n\nBased on the context above, please explain the following highlighted portion specifically for a {request.complexity}. If it's a single word, define it simply. Avoid jargon:\n\"{request.highlighted_text}\""
            messages.append({"role": "user", "content": initial_prompt})
            
            for msg in request.chat_history:
                role = "assistant" if msg.role == "model" else "user"
                messages.append({"role": role, "content": msg.content})
                
            if request.new_question:
                messages.append({"role": "user", "content": request.new_question})
                
            response = openai_client.chat.completions.create(
                model="gpt-4o-mini",
                messages=messages,
                temperature=0.3,
                stream=True
            )
            for chunk in response:
                if chunk.choices and chunk.choices[0].delta.content:
                    yield chunk.choices[0].delta.content
        except Exception as e:
            error_msg = str(e)
            print(f"Error calling openai stream: {error_msg}")
            if "429" in error_msg or "RESOURCE_EXHAUSTED" in error_msg or "RateLimitError" in error_msg:
                yield "\n\n**Too many rapid requests!** You're decoding jargon faster than the engine can process. Please wait 15 seconds and try again."
            elif "503" in error_msg or "UNAVAILABLE" in error_msg:
                yield "\n\n**API Busy:** The model is currently experiencing high demand. Please try again in a few seconds."
            else:
                yield f"\n\n**Error:** Failed to simplify text."

    return StreamingResponse(generate(), media_type="text/plain")

@app.post("/upload-document")
async def upload_document(file: UploadFile = File(...)):
    filename = file.filename.lower()
    if not (filename.endswith(".pdf") or filename.endswith(".docx")):
        raise HTTPException(status_code=400, detail="Only PDF and DOCX files are supported.")
    
    try:
        content = await file.read()
        extracted_text = ""
        
        if filename.endswith(".pdf"):
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(content))
            for page in pdf_reader.pages:
                text = page.extract_text()
                if text:
                    extracted_text += text + "\n\n"
        elif filename.endswith(".docx"):
            doc = docx.Document(io.BytesIO(content))
            for para in doc.paragraphs:
                if para.text.strip():
                    extracted_text += para.text + "\n\n"
                
        return {"text": extracted_text.strip()}
    except Exception as e:
        print(f"Error parsing document: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Failed to parse document: {str(e)}")

@app.post("/export-docx")
async def export_docx(request: ExportRequest):
    try:
        doc = docx.Document()
        doc.add_heading("Layman Lens Document", level=1)
        
        for paragraph in request.text.split('\n'):
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip())
                
        # Save to BytesIO
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        
        return Response(
            content=file_stream.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": 'attachment; filename="simplified_document.docx"'}
        )
    except Exception as e:
        print(f"Error exporting docx: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Failed to export document: {str(e)}")

# Mount static files (HTML, JS, CSS) at the root
app.mount("/", StaticFiles(directory=".", html=True), name="static")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=True)
